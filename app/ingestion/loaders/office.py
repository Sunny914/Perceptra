import logfire
from docx import Document
from pptx import Presentation


def parse_office(file_path: str) -> str:
    """Extract text from .docx and .pptx using the libraries already declared for ingestion.

    python-docx and python-pptx cover the formats this loader accepts. The previous
    unstructured import was not declared in pyproject.toml, so Office parsing failed
    in the Docker image.
    """
    with logfire.span("Office Document Parsing", filename=file_path):
        ext = file_path.lower().rsplit(".", 1)[-1]
        try:
            if ext == "docx":
                document = Document(file_path)
                full_text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
            elif ext == "pptx":
                presentation = Presentation(file_path)
                texts = []
                for slide in presentation.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            texts.append(shape.text.strip())
                full_text = "\n".join(texts)
            else:
                raise ValueError(f"Unsupported office extension: {ext}")

            if not full_text.strip():
                logfire.warning(f"Office parser returned empty text for {file_path}")
            else:
                logfire.info(f"Parsed {len(full_text)} characters from {file_path}")
            return full_text
        except Exception as e:
            logfire.error(f"Office parse failed: {e}")
            raise
